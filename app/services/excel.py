"""app/services/excel.py — geração de planilhas .xlsx para exportação."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_workbook(sheets: dict) -> io.BytesIO:
    """
    sheets = { "NomeAba": {"headers": [...], "rows": [[...], ...]} }
    Retorna um BytesIO pronto para send_file.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        headers = data.get("headers", [])
        rows = data.get("rows", [])
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        # auto width
        for c in range(1, len(headers) + 1):
            width = len(str(headers[c - 1])) + 2
            for row in rows:
                if c - 1 < len(row):
                    width = max(width, len(str(row[c - 1])) + 2)
            ws.column_dimensions[get_column_letter(c)].width = min(width, 55)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
